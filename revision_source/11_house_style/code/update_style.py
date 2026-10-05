"""Bounded JHG house-style edits; no numerical or analytical changes."""
from pathlib import Path
import re,json,shutil,hashlib,difflib
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from xml.etree import ElementTree as ET
import cairosvg

B=Path(__file__).resolve().parents[1];ROOT=B.parents[1];D=ROOT/'submission/JHG'
BASE=B/'baseline';RES=B/'results';RES.mkdir(parents=True,exist_ok=True)
assert not BASE.exists(),'Do not overwrite the reviewed baseline'
shutil.copytree(D,BASE)
(RES/'prior_zip_sha256.txt').write_text(hashlib.sha256((ROOT/'submission/JHG_Review_Package.zip').read_bytes()).hexdigest()+'\n')
comma=re.compile(r'(?<!\d)\d{1,3}(?:,\d{3})+(?!\d)')
def number_style(s):
    def change(m):
        raw=m[0].replace(',','')
        return raw if len(raw)==4 else m[0].replace(',','\u00a0')
    return comma.sub(change,s)
replacements={
 'Main_Manuscript':[
  ('European LDetect and frozen ASN LDetect boundaries','European LDetect and frozen Asian (ASN) LDetect boundaries'),
  ('its annotated transcription start site or any position','its annotated transcription start site (TSS) or any position'),
  ('figure production and manuscript drafting','generation and refinement of plotting code, and manuscript drafting'),
  ('confirm these actual uses and approve responsibility for the code, numerical results, citations and interpretation.','confirm these actual uses, whether any generative image model was used, and researcher verification and responsibility for the code, figures, numerical results, citations and interpretation.'),
  ('left extensive representation of seed-gene-body inputs among recurrent candidates','left representation of seed-gene-body inputs among recurrently supported candidates'),
  ('when its TSS or any frozen mapped position','when its transcription start site (TSS) or any frozen mapped position'),
  ('EUR denotes European and EAS East Asian assessment.','EUR denotes European and EAS East Asian assessment; ASN denotes the Asian LDetect block resource.')],
 'Supplementary_Methods':[
  ('LDetect EUR/ASN intervals supply block boundaries','European (EUR) and Asian (ASN) LDetect intervals supply block boundaries'),
  ('Across the original 283 candidates, TSS blocks','Across the original 283 candidates, transcription start site (TSS) blocks')],
 'Cover_Letter':[
  ('It shows why excluding known genes can leave their regional evidence distributed across other candidates','It shows that excluding known genes can leave their regional evidence represented across other candidates')],
 'Table_2':[]}
def change(text,name):
    for old,new in replacements[name]:text=text.replace(old,new)
    return number_style(text)
def replace_preserving_runs(p,old,new):
    for tag,a,b,c,d in reversed(difflib.SequenceMatcher(a=old,b=new,autojunk=False).get_opcodes()):
        if tag=='equal':continue
        positions=[];offset=0
        for run in p.runs:
            positions.append((run,offset,offset+len(run.text)));offset+=len(run.text)
        hit=False
        for run,start,end in positions:
            if a==b:
                if not hit and start<=a<=end:
                    run.text=run.text[:a-start]+new[c:d]+run.text[a-start:];hit=True
            elif start<b and end>a:
                left=run.text[:max(0,a-start)];right=run.text[max(0,b-start):] if b<end else ''
                run.text=left+(new[c:d] if not hit else '')+right;hit=True
        assert hit
    assert p.text==new
edit_log={}
for name in replacements:
    md=D/f'{name}.md'
    if md.exists():
        text=md.read_text()
        for old,new in replacements[name]:assert old in text,(name,old)
        md.write_text(change(text,name))
    doc=Document(D/f'{name}.docx');changed=[]
    paragraphs=list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:paragraphs.extend(cell.paragraphs)
    seen=set()
    for p in paragraphs:
        if p._p in seen:continue
        seen.add(p._p)
        old=p.text;new=change(old,name)
        if old==new:continue
        changed.append({'old':old,'new':new})
        replace_preserving_runs(p,old,new)
    assert changed,name
    doc.save(D/f'{name}.docx');edit_log[name]=changed
w=load_workbook(D/'Table_2.xlsx');cells=[]
for s in w:
    for row in s:
        for c in row:
            if isinstance(c.value,str) and number_style(c.value)!=c.value:
                cells.append([s.title,c.coordinate,c.value,number_style(c.value)]);c.value=number_style(c.value)
w.save(D/'Table_2.xlsx');edit_log['Table_2_XLSX']=cells

# Modify only visible text in Figure 1, not SVG coordinates or shapes.
p=D/'Main_Figures/Figure_1.svg';svg=p.read_text()
for old,new in [('1,652','1652'),('1,259','1259'),('1,465','1465'),('3,224','3224')]:
    assert svg.count(old)==1;svg=svg.replace(old,new)
p.write_text(svg)
before=ET.parse(BASE/'Main_Figures/Figure_1.svg').getroot();after=ET.parse(p).getroot()
assert len(list(before.iter()))==len(list(after.iter()))
for a,b in zip(before.iter(),after.iter()):
    assert a.tag==b.tag and a.attrib==b.attrib and a.tail==b.tail
    assert b.text==(number_style(a.text) if a.text else a.text)
width,height=Image.open(BASE/'Main_Figures/Figure_1.png').size
cairosvg.svg2pdf(bytestring=svg.encode(),write_to=str(p.with_suffix('.pdf')))
cairosvg.svg2png(bytestring=svg.encode(),write_to=str(p.with_suffix('.png')),output_width=width,output_height=height)
edit_log['Figure_1']='Four comma separators removed in text nodes; all SVG geometry and attributes unchanged'
edit_log['Figure_3']='Already comma-free in SVG/PDF/PNG; unchanged'
(RES/'edits.json').write_text(json.dumps(edit_log,indent=2,ensure_ascii=False)+'\n')
print('House-style edits ready; Figure 3 and all supplementary workbooks unchanged.')
