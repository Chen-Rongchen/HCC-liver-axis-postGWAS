"""Build coherent corrected tables and main figures; no new scientific models."""
from pathlib import Path
import json,shutil,sys
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill
from openpyxl.utils import get_column_letter
BASE=Path(__file__).resolve().parents[1];ROOT=BASE.parents[1];D=BASE/'stage/JHG';R=BASE/'results/models';T=D/'Supplementary_Source_TSV'
AUDIT=ROOT/'analysis_signal_resolution_c/08_genomewide_input_audit/results'
A=ROOT/'analysis_hbsn_restructure/35_seed_region_masking/out_03_membership_footprint'
def read(p):return pd.read_csv(p,sep='\t',float_precision='round_trip')
def write(x,p):x.to_csv(p,sep='\t',index=False,float_format='%.17g')
for src,dst in [('model_estimates.tsv','S09_model_estimates.tsv'),('overall_threshold_context.tsv','S10_overall_threshold_context.tsv'),('block_influence_summary.tsv','S11_block_influence_summary.tsv')]:shutil.copy2(R/src,T/dst)
models=read(AUDIT/'M2_impact.tsv');enrich=read(AUDIT/'enrichment_impact.tsv')
rule=models.merge(enrich,on=['ancestry','version','part'],validate='one_to_one')
rule['high_n']=np.where(rule.part.eq('original'),443,214);rule['low_n']=rule.n-rule.high_n
rule['analysis_role']=rule.version.map({'frozen':'historical_baseline','collision_repaired':'current_manuscript','roundtrip_exclusion_sensitivity':'conservative_mapping_sensitivity'})
write(rule,T/'S25_membership_sensitivity.tsv')
classes=read(A/'gene_mask_class_A_gene_body.tsv').rename(columns={'gene':'gene_symbol'})
classes['excluded']=~classes.mask_class.isin(['retained','retained_tss_unresolved'])
write(classes,T/'S26_whole_universe_mask.tsv')
candidate=read(A/'A_candidate_membership.tsv');ann=read(T/'S13_gene_annotation_audit.tsv')
candidate=candidate.merge(ann[['gene_symbol','audit_final_EUR_block']],on='gene_symbol',validate='one_to_one')
candidate['retained_by_rule_A']=candidate['class'].isin(['retained','retained_tss_unresolved'])
candidate['reporting_scope']='original_EUR_top5_B_ge4_selected_set'
write(candidate,T/'S27_candidate_seed_input_footprint.tsv')
conflicts=[]
for anc,acc in [('EUR','GCST90860790'),('EAS','GCST90860791')]:
    x=read(AUDIT/acc/'conflicting_source_records.tsv');x.insert(0,'ancestry',anc);x.insert(1,'accession',acc);x['corrected_action']='all_records_for_target_key_excluded';conflicts.append(x)
write(pd.concat(conflicts),T/'S28_mapped_key_conflict_records.tsv')
x=read(ROOT/'analysis_signal_resolution_c/07_code_corrections/results/molecular28_corrected_summary.tsv');x['current_role']='historical_only_no_reliable_signal_sharing_inference'
write(x,T/'S29_historical_molecular_QTL_status.tsv')
for number in [4,5,6,7,19,20,23]:
    p=next(T.glob(f'S{number:02d}_*.tsv'));x=read(p)
    if number in [5,6,20]:x['current_inference_status']='historical_registry_only_effect_encoding_N_and_LD_not_certified'
    elif number==4:
        history=x.parameter.isin(['coloc_coverage','frequency_gate','susie','coloc_prior','pair_LD_gate','RSS_N','coloc_status'])
        x['current_role']=np.where(history,'historical_parameters_not_validated_for_inference','retained_parameter_contract_historical_source_link')
        x.loc[x.parameter.eq('coloc_status'),'implemented_value']='Historical software-output status only; no reliable signal-sharing inference'
    elif number==7:x['source_version']='historical_environment; current_revision_environment_in_revision_provenance'
    elif number==19:x['historical_evaluable_fields_note']='n_evaluable* fields describe prior software stage only'
    elif number==23:x['historical_evaluable_fields_note']='Metrics containing evaluable describe prior software stage only'
    write(x,p)
# Workbook is rebuilt from the exact TSVs, preserving numeric types.
w=Workbook();w.remove(w.active)
for path in sorted(T.glob('S*.tsv')):
    number=int(path.name[1:3]);data=read(path);sheet=w.create_sheet(f'S{number}');sheet.append(list(data.columns))
    for row in data.itertuples(index=False,name=None):sheet.append([None if pd.isna(v) else str(v) if isinstance(v,(int,float)) and not np.isfinite(v) else v for v in row])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for c in sheet[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='286782');c.alignment=Alignment(wrap_text=True,vertical='top')
    sheet.row_dimensions[1].height=45
    for i,c in enumerate(data.columns,1):sheet.column_dimensions[get_column_letter(i)].width=min(55,max(16,len(c)*.8))
w.save(D/'Supplementary_Tables.xlsx')

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'svg.fonttype':'none','pdf.fonttype':42,'axes.spines.top':False,'axes.spines.right':False})
F=D/'Main_Figures';F.mkdir(exist_ok=True)
def save(fig,n):
    for ext in ['pdf','svg','png']:fig.savefig(F/f'Figure_{n}.{ext}',dpi=600,facecolor='white',bbox_inches='tight')
    plt.close(fig)
fig,ax=plt.subplots(figsize=(7.2,5.8));ax.set_axis_off();ax.set_xlim(0,1);ax.set_ylim(0,1)
def box(x,y,width,height,text,color='#EDF4F7'):
    ax.add_patch(FancyBboxPatch((x,y),width,height,boxstyle='round,pad=0.008,rounding_size=0.01',linewidth=.7,edgecolor='#70929D',facecolor=color))
    ax.text(x+width/2,y+height/2,text,ha='center',va='center',fontsize=9,linespacing=1.5)
ax.text(0,1,'a  Source-mapped gene support',fontweight='bold',va='top',fontsize=11)
box(.02,.77,.28,.14,'11 non-HCC liver traits\nFixed source cis records')
box(.37,.77,.28,.14,'Aggregated input ranks\nTop 5% defines support')
box(.72,.77,.26,.14,'56 seed identities\nexcluded')
for x1,x2 in [(.30,.37),(.65,.72)]:ax.annotate('',xy=(x2,.84),xytext=(x1,.84),arrowprops=dict(arrowstyle='->',color='#286782'))
ax.text(.5,.72,'Mapped input ≠ significant eQTL link or effector-gene assignment',ha='center',fontsize=9,color='#485C64')
ax.text(0,.65,'b  Distinct analysis populations',fontweight='bold',va='top',fontsize=11)
box(.02,.40,.46,.18,'Conditional breadth comparison\n1,652 genes → 1,259 after masking\n≥4 versus 1–3 supporting traits')
box(.53,.40,.45,.18,'Selected input audit\n283 genes · 1,465 support events\n3,224 within-gene trait pairs')
ax.text(.5,.35,'HCC association is assessed separately in European and East Asian data',ha='center',fontsize=9,color='#485C64')
ax.text(0,.28,'c  Seed-gene-body membership mask (Rule A)',fontweight='bold',va='top',fontsize=11)
box(.02,.065,.96,.15,'Exclude the entire gene if its TSS OR any frozen mapped input overlaps a seed body.\nRetain the existing support flags, scores and global HCC ranks.\nCompare populations; do not attribute the change to a causal input contribution.','#F2F5EF')
save(fig,1)

current=rule[rule.version.eq('collision_repaired')].set_index(['ancestry','part'])
fig,ax=plt.subplots(figsize=(7.2,3.9));ys=[3.4,2.55,1.25,.4];rows=[('EUR','original'),('EUR','retained'),('EAS','original'),('EAS','retained')]
for y,(anc,part) in zip(ys,rows):
    r=current.loc[(anc,part)];color='#286782' if anc=='EUR' else '#AD633C'
    ax.errorbar(r.estimate,y,xerr=[[r.estimate-r.ci_low],[r.ci_high-r.estimate]],fmt='o' if part=='original' else 's',color=color,markersize=6,capsize=3,lw=1.5)
    ax.text(1.52,y,f'{r.estimate:.3f} [{r.ci_low:.3f}, {r.ci_high:.3f}]',va='center',fontsize=9)
ax.axvline(0,color='#777777',ls='--',lw=.8)
ax.set_yticks(ys,labels=['EUR · original','EUR · masked','EAS · original','EAS · masked']);ax.set_ylim(-.25,4.15)
ax.set_xlim(-.9,2.9);ax.set_xticks([-.5,0,.5,1]);ax.spines['left'].set_visible(False);ax.tick_params(axis='y',length=0)
ax.spines['bottom'].set_bounds(-.8,1.4);ax.set_xlabel('Conditional breadth contrast (rank-evidence units)',loc='left')
ax.text(1.52,3.95,'Estimate [95% CI]',fontweight='bold',fontsize=9)
ax.set_title('M2 before and after seed-gene-body membership masking',loc='left',fontsize=11,fontweight='bold',pad=16)
fig.tight_layout();save(fig,2)
# Original Figure 3 and its source tables retain identical bytes.
source=D/'Figure_Source_Data'
for p in list(source.glob('Fig1*'))+list(source.glob('Fig2*')):p.unlink()
write(rule[rule.version.eq('collision_repaired')],source/'Fig2_M2_membership_sensitivity.tsv')
write(pd.DataFrame([dict(population='original_model',n=1652,high=443,low=1209),dict(population='masked_model',n=1259,high=214,low=1045),dict(population='selected_input_audit',n=283,high=283,low=0)]),source/'Fig1_analysis_populations.tsv')
if (D/'Main_Figure_Panels').exists():shutil.rmtree(D/'Main_Figure_Panels')

# Independent editable main tables; DOCX is the chosen upload format.
from docx import Document
from docx.shared import Inches,Pt
from openpyxl.worksheet.page import PageMargins
main={}
data=[]
for anc,part in rows:
    r=current.loc[(anc,part)]
    data.append([anc,'Original' if part=='original' else 'Masked',int(r.nominee_top5),int(r.nominee_total),int(r.background_top5),int(r.background_total),float(r.OR)])
main['1']=dict(title='Table 1. Descriptive HCC top-5% enrichment before and after membership masking',headers=['Ancestry','Population','Supported: top 5%','Supported: total','Background: top 5%','Background: total','Odds ratio'],rows=data,note='EUR, European; EAS, East Asian; HCC, hepatocellular carcinoma. Both groups exclude 56 fixed seed identities. Masked populations additionally exclude genes whose TSS or any frozen mapped input overlaps a seed gene body. HCC ranks remain fixed. Odds ratios are descriptive and do not adjust for regional dependence; Fisher P values and input versions are in Table S25.')
main['2']=dict(title='Table 2. Gene identities and represented inputs in the selected candidate set',headers=['Measurement','Count','Unit / interpretation'],rows=[['Original selected set',283,'Non-seed genes with ≥4 supporting traits and EUR HCC top-5% evidence'],['Minimum input in a seed body',157,'Candidate genes; 607 of 1,465 supporting events'],['Excluded by the full membership mask',172,'Candidate genes with TSS or any frozen-input overlap'],['Retained after masking',111,'Candidate genes; not a new independently selected validation set'],['Retained with chr6 TSS-block assignment',107,'Genes assigned within chr6:23.9–29.7 Mb'],['Retained without TSS-block assignment',3,'Genes whose mapped inputs also involve that chr6 region'],['Other retained gene',1,'DUSP10; three support events use MTARC1-associated rs2642438 as minimum input']],note='Minimum-input overlap and the whole-gene membership mask use different criteria. Gene counts do not measure independent genetic signals or identify causal effector genes. Original support events and pairs are in Tables S14–S15; complete candidate membership is in Table S27.')
for key,t in main.items():
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.6)
    doc.add_heading(t['title'],1);tab=doc.add_table(rows=1,cols=len(t['headers']));tab.style='Table Grid'
    for cell,val in zip(tab.rows[0].cells,t['headers']):cell.text=val
    for row in t['rows']:
        for j,(cell,val) in enumerate(zip(tab.add_row().cells,row)):
            cell.text=f'{val:.2f}' if isinstance(val,float) else str(val)
    widths=[.7,.9,1.3,1.3,1.5,1.5,1.0] if key=='1' else [3.1,.7,6.3]
    tab.autofit=False
    for row in tab.rows:
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
                for run in p.runs:run.font.size=Pt(10);run.font.name='Arial'
    doc.add_paragraph(t['note']);doc.save(D/f'Table_{key}.docx')
    wb=Workbook();ws=wb.active;ws.title=f'Table {key}';n=len(t['headers'])
    ws.append([t['title']]);ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n);ws.append(t['headers'])
    for row in t['rows']:ws.append(row)
    note_row=len(t['rows'])+4;ws.cell(note_row,1,t['note']);ws.merge_cells(start_row=note_row,start_column=1,end_row=note_row,end_column=n)
    ws.row_dimensions[1].height=35;ws.row_dimensions[2].height=45;ws.row_dimensions[note_row].height=80
    for j,width in enumerate(([15,18,22,22,25,25,18] if key=='1' else [52,12,85]),1):ws.column_dimensions[get_column_letter(j)].width=width
    for row in ws:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='center');cell.font=Font(name='Arial',size=11)
    for i in range(3,len(t['rows'])+3):ws.row_dimensions[i].height=48 if key=='2' else 30
    for cell in ws[2]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='286782')
    if key=='1':
        for i in range(3,7):ws.cell(i,7).number_format='0.00'
    ws.print_area=f'A1:{get_column_letter(n)}{note_row}';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1;ws.page_margins=PageMargins(left=.25,right=.25,top=.4,bottom=.4)
    ws.freeze_panes='A3';wb.save(D/f'Table_{key}.xlsx')
for p in D.glob('Main_Tables.*'):p.unlink()
(BASE/'results/main_table_values.json').write_text(json.dumps(main,indent=2,ensure_ascii=False)+'\n')
print('Assets synchronized: S1–S29, Tables 1–2, Figures 1–3')
