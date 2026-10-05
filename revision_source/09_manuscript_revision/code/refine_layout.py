from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from openpyxl import load_workbook
from openpyxl.styles import Alignment,PatternFill
BASE=Path(__file__).resolve().parents[1];D=BASE/'stage/JHG'
for n in [1,2]:
    p=D/f'Table_{n}.docx';d=Document(p);d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10)
    widths=[.7,.9,1.3,1.3,1.5,1.5,1.0] if n==1 else [3.1,.7,6.3]
    for col,width in zip(d.tables[0].columns,widths):col.width=Inches(width)
    if n==2:
        for row in d.tables[0].rows:
            for para in row.cells[1].paragraphs:para.alignment=1
    d.save(p)
    if n==2:
        w=load_workbook(D/'Table_2.xlsx');s=w.active
        for i in range(3,10):
            s.cell(i,2).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            s.cell(i,3).alignment=Alignment(horizontal='left',vertical='center',wrap_text=True,indent=1)
            if i%2:
                for c in s[i]:c.fill=PatternFill('solid',fgColor='F1F5F7')
        w.save(D/'Table_2.xlsx')
p=D/'Cover_Letter.docx';d=Document(p)
for para in d.paragraphs:para.paragraph_format.line_spacing=1.15
d.save(p)
